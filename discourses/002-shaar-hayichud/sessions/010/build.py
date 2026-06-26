"""Build a Hebrew RTL Word document from this session's transcript.md.

Uses python-docx primitives directly. Does not import any other build script
in the repo (per CLAUDE.md). Footnotes are written via direct OOXML
construction because python-docx 1.2.0 has no high-level footnote API.
"""

from __future__ import annotations

import re
from html import escape as xml_escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.parts.story import StoryPart
from docx.shared import Cm, Pt

# ----- paths -----------------------------------------------------------------
HERE = Path(__file__).resolve().parent
DISCOURSE_ROOT = HERE.parent.parent  # discourses/002-shaar-hayichud
TRANSCRIPT = HERE / "transcript.md"
OUTPUT_DIR = DISCOURSE_ROOT / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUTPUT_DIR / "shaar-hayichud-perek-alef-001-he.docx"

HEBREW_FONT = "David"
FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)


# ----- low-level XML helpers -------------------------------------------------
def _set_rtl_run(run) -> None:
    rpr = run._element.get_or_add_rPr()
    for tag in ("rtl", "cs"):
        if rpr.find(qn(f"w:{tag}")) is None:
            el = OxmlElement(f"w:{tag}")
            el.set(qn("w:val"), "1")
            rpr.append(el)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), HEBREW_FONT)
    rfonts.set(qn("w:ascii"), HEBREW_FONT)
    rfonts.set(qn("w:hAnsi"), HEBREW_FONT)


def _set_rtl_paragraph(paragraph) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.append(bidi)


def _set_rtl_section(section) -> None:
    sectPr = section._sectPr
    if sectPr.find(qn("w:bidi")) is None:
        sectPr.append(OxmlElement("w:bidi"))


def _set_line_spacing(paragraph, multiple: float) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing")) or OxmlElement("w:spacing")
    if spacing.getparent() is None:
        pPr.append(spacing)
    spacing.set(qn("w:line"), str(int(multiple * 240)))
    spacing.set(qn("w:lineRule"), "auto")


def _set_first_line_indent(paragraph, cm: float) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind")) or OxmlElement("w:ind")
    if ind.getparent() is None:
        pPr.append(ind)
    ind.set(qn("w:firstLine"), str(int(cm * 567)))


def _set_block_indent(paragraph, left_cm: float, right_cm: float) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind")) or OxmlElement("w:ind")
    if ind.getparent() is None:
        pPr.append(ind)
    ind.set(qn("w:left"), str(int(left_cm * 567)))
    ind.set(qn("w:right"), str(int(right_cm * 567)))


# ----- footnote helpers ------------------------------------------------------
def _build_footnotes_xml(definitions: list[tuple[int, str]]) -> bytes:
    parts: list[str] = []
    parts.append('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')
    parts.append(
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    )
    parts.append(
        '<w:footnote w:type="separator" w:id="-1">'
        "<w:p><w:r><w:separator/></w:r></w:p>"
        "</w:footnote>"
    )
    parts.append(
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        "<w:p><w:r><w:continuationSeparator/></w:r></w:p>"
        "</w:footnote>"
    )
    for fn_id, text in definitions:
        safe = xml_escape(text, quote=True)
        parts.append(
            f'<w:footnote w:id="{fn_id}">'
            "<w:p>"
            '<w:pPr><w:bidi w:val="1"/></w:pPr>'
            "<w:r>"
            '<w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
            "<w:footnoteRef/>"
            "</w:r>"
            '<w:r><w:t xml:space="preserve"> </w:t></w:r>'
            "<w:r>"
            "<w:rPr>"
            f'<w:rFonts w:ascii="{HEBREW_FONT}" w:hAnsi="{HEBREW_FONT}" w:cs="{HEBREW_FONT}"/>'
            '<w:rtl w:val="1"/>'
            '<w:cs w:val="1"/>'
            '<w:sz w:val="22"/>'
            '<w:szCs w:val="22"/>'
            "</w:rPr>"
            f"<w:t>{safe}</w:t>"
            "</w:r>"
            "</w:p>"
            "</w:footnote>"
        )
    parts.append("</w:footnotes>")
    return "".join(parts).encode("utf-8")


def _attach_footnotes_part(doc, definitions: list[tuple[int, str]]) -> None:
    if not definitions:
        return
    blob = _build_footnotes_xml(definitions)
    element = parse_xml(blob)
    partname = PackURI("/word/footnotes.xml")
    footnotes_part = StoryPart(partname, FOOTNOTES_CT, element, doc.part.package)
    doc.part.relate_to(footnotes_part, RELATIONSHIP_TYPE.FOOTNOTES)


def _add_footnote_reference(paragraph, footnote_id: int) -> None:
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "FootnoteReference")
    rPr.append(rStyle)
    r.append(rPr)
    fnref = OxmlElement("w:footnoteReference")
    fnref.set(qn("w:id"), str(footnote_id))
    r.append(fnref)
    paragraph._element.append(r)


# ----- transcript parsing ----------------------------------------------------
FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n", re.DOTALL)
UNDERLINE_RE = re.compile(r"\[([^\[\]]+)\]\{\.underline\}")
HTML_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
FOOTNOTE_REF_RE = re.compile(r"\[\^([A-Za-z0-9_\-]+)\]")
FOOTNOTE_DEF_RE = re.compile(r"^\[\^([A-Za-z0-9_\-]+)\]:\s*(.+)$", re.MULTILINE)


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    m = FRONTMATTER_RE.match(text)
    if not m:
        return {}, text
    fm_block, body = m.group(1), text[m.end() :]
    fm: dict[str, str] = {}
    for line in fm_block.splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        key = key.strip()
        value = value.strip()
        if value.startswith('"') and value.endswith('"'):
            value = value[1:-1]
        fm[key] = value
    return fm, body


def extract_footnote_definitions(body: str) -> tuple[dict[str, str], str]:
    defs: dict[str, str] = {}
    matches = list(FOOTNOTE_DEF_RE.finditer(body))
    if not matches:
        return defs, body
    earliest = matches[0].start()
    defs_block = body[earliest:]
    for m in FOOTNOTE_DEF_RE.finditer(defs_block):
        defs[m.group(1)] = m.group(2).strip()
    return defs, body[:earliest].rstrip() + "\n"


def parse_blocks(body: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in re.split(r"\n\s*\n", body.strip()):
        text = raw.strip()
        if not text:
            continue
        if text.startswith("## "):
            blocks.append(("h2", text[3:].strip()))
        elif text.startswith("# "):
            blocks.append(("h1", text[2:].strip()))
        elif text.startswith("> "):
            lines = [
                ln[2:] if ln.startswith("> ") else ln.lstrip("> ")
                for ln in text.splitlines()
            ]
            blocks.append(("quote", " ".join(lines).strip()))
        else:
            cleaned = HTML_COMMENT_RE.sub("", text)
            cleaned = re.sub(r"\s+", " ", cleaned).strip()
            if cleaned:
                blocks.append(("para", cleaned))
    return blocks


# ----- emitters --------------------------------------------------------------
def emit_runs(paragraph, text: str, footnote_id_map: dict[str, int]) -> list:
    """Emit runs for `text`, handling [^N] footnotes and [..]{.underline}.

    Returns the list of created runs (excluding footnote-reference XML runs)
    so the caller can apply font/size styling.
    """
    markers: list[tuple[int, int, str, str]] = []
    for m in FOOTNOTE_REF_RE.finditer(text):
        markers.append((m.start(), m.end(), "footnote", m.group(1)))
    for m in UNDERLINE_RE.finditer(text):
        markers.append((m.start(), m.end(), "underline", m.group(1)))
    markers.sort(key=lambda x: x[0])

    runs: list = []
    pos = 0
    for start, end, kind, content in markers:
        if start > pos:
            run = paragraph.add_run(text[pos:start])
            _set_rtl_run(run)
            runs.append(run)
        if kind == "underline":
            run = paragraph.add_run(content)
            run.underline = True
            _set_rtl_run(run)
            runs.append(run)
        elif kind == "footnote":
            fn_id = footnote_id_map.get(content)
            if fn_id is not None:
                _add_footnote_reference(paragraph, fn_id)
        pos = end
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_rtl_run(run)
        runs.append(run)
    return runs


def add_title(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl_paragraph(p)
    _set_line_spacing(p, 1.4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(30)
    run.font.name = HEBREW_FONT
    _set_rtl_run(run)


def add_subtitle(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl_paragraph(p)
    _set_line_spacing(p, 1.4)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(14)
    run.font.name = HEBREW_FONT
    _set_rtl_run(run)


def add_heading(doc, text: str, *, level: int = 2) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl_paragraph(p)
    _set_line_spacing(p, 1.5)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20 if level == 1 else 16)
    run.font.name = HEBREW_FONT
    _set_rtl_run(run)


def add_body(doc, text: str, footnote_id_map: dict[str, int]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_rtl_paragraph(p)
    _set_line_spacing(p, 1.7)
    _set_first_line_indent(p, 0.8)
    runs = emit_runs(p, text, footnote_id_map)
    for run in runs:
        run.font.size = Pt(13)
        run.font.name = HEBREW_FONT


def add_quote(doc, text: str, footnote_id_map: dict[str, int]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_rtl_paragraph(p)
    _set_line_spacing(p, 1.5)
    _set_block_indent(p, left_cm=1.0, right_cm=1.0)
    runs = emit_runs(p, text, footnote_id_map)
    for run in runs:
        run.font.size = Pt(13)
        run.font.name = HEBREW_FONT
        run.italic = True


# ----- main ------------------------------------------------------------------
def build() -> Path:
    raw = TRANSCRIPT.read_text(encoding="utf-8")
    fm, body = parse_frontmatter(raw)
    footnote_defs, body_no_defs = extract_footnote_definitions(body)
    blocks = parse_blocks(body_no_defs)

    # Map markdown footnote labels to numeric IDs (Word footnotes start at 1).
    # Walk the body in order so footnote numbering matches first appearance.
    footnote_id_map: dict[str, int] = {}
    next_id = 1
    seen_in_body = []
    for _, text in blocks:
        for m in FOOTNOTE_REF_RE.finditer(text):
            label = m.group(1)
            if label not in footnote_id_map and label in footnote_defs:
                footnote_id_map[label] = next_id
                seen_in_body.append((next_id, footnote_defs[label]))
                next_id += 1

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = HEBREW_FONT
    style.font.size = Pt(13)

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    _set_rtl_section(section)

    add_title(doc, fm.get("title", "שער היחוד"))
    if fm.get("subtitle"):
        add_subtitle(doc, fm["subtitle"])

    for kind, text in blocks:
        if kind == "h1":
            add_heading(doc, text, level=1)
        elif kind == "h2":
            add_heading(doc, text, level=2)
        elif kind == "quote":
            add_quote(doc, text, footnote_id_map)
        else:
            add_body(doc, text, footnote_id_map)

    _attach_footnotes_part(doc, seen_in_body)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(f"wrote {out}")
    print(f"size: {out.stat().st_size:,} bytes")
